from __future__ import annotations

import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


app = FastAPI(
    title="POP Audit Report Service",
    version="2.0.0",
)


REPORT_DIR = Path(
    os.getenv(
        "REPORT_DIR",
        "/opt/pop-audit/reports",
    )
)

REPORT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


STATUS_PRIORITY = {
    "CRITICAL": 5,
    "WARNING": 4,
    "UNKNOWN": 3,
    "NOT_AVAILABLE": 2,
    "PASS": 1,
}


STATUS_FILL = {
    "PASS": "C6EFCE",
    "WARNING": "FFEB9C",
    "CRITICAL": "FFC7CE",
    "UNKNOWN": "D9E1F2",
    "NOT_AVAILABLE": "E7E6E6",
}


HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
TITLE_FILL = "17365D"


thin_gray = Side(
    style="thin",
    color="D9E1F2",
)

BORDER = Border(
    left=thin_gray,
    right=thin_gray,
    top=thin_gray,
    bottom=thin_gray,
)


# ============================================================
# Status helpers
# ============================================================

def aggregate_status(
    statuses: list[str],
) -> str:

    if not statuses:
        return "UNKNOWN"

    return max(
        statuses,
        key=lambda value: STATUS_PRIORITY.get(
            value,
            3,
        ),
    )


def service_status(
    result: Any,
) -> str:

    if not isinstance(
        result,
        dict,
    ):
        return "UNKNOWN"

    status = result.get("status")

    if status in STATUS_PRIORITY:
        return status

    if status == "COMPLETED":

        data = result.get(
            "data"
        )

        if isinstance(data, dict):

            data_status = data.get(
                "status"
            )

            if data_status in STATUS_PRIORITY:
                return data_status

        return "PASS"

    if status == "FAILED":
        return "CRITICAL"

    return "UNKNOWN"


def find_data(
    result: dict[str, Any],
) -> dict[str, Any]:

    data = result.get(
        "data"
    )

    if isinstance(data, dict):
        return data

    return result


# ============================================================
# Excel helpers
# ============================================================

def style_header(
    ws,
    row: int,
) -> None:

    for cell in ws[row]:

        if cell.value is None:
            continue

        cell.fill = PatternFill(
            "solid",
            fgColor=HEADER_FILL,
        )

        cell.font = Font(
            bold=True,
            color=HEADER_FONT,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.border = BORDER


def style_title(
    ws,
    cell_range: str,
) -> None:

    for row in ws[cell_range]:

        for cell in row:

            cell.fill = PatternFill(
                "solid",
                fgColor=TITLE_FILL,
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF",
                size=14,
            )

            cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
            )


def style_body(
    ws,
    min_row: int,
    max_row: int,
    max_col: int,
) -> None:

    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=1,
        max_col=max_col,
    ):

        for cell in row:

            cell.border = BORDER

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


def style_status_cell(
    cell,
) -> None:

    status = str(
        cell.value or ""
    ).upper()

    color = STATUS_FILL.get(
        status
    )

    if not color:
        return

    cell.fill = PatternFill(
        "solid",
        fgColor=color,
    )

    cell.font = Font(
        bold=True,
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )


def auto_width(
    ws,
    max_width: int = 38,
) -> None:

    for column_cells in ws.columns:

        length = 0

        for cell in column_cells:

            value = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            if "\n" in value:
                value = max(
                    value.split("\n"),
                    key=len,
                )

            length = max(
                length,
                len(value),
            )

        width = min(
            max(length + 2, 10),
            max_width,
        )

        ws.column_dimensions[
            get_column_letter(
                column_cells[0].column
            )
        ].width = width


def add_kpi(
    ws,
    row: int,
    label: str,
    value: Any,
) -> None:

    ws.cell(
        row=row,
        column=1,
        value=label,
    )

    ws.cell(
        row=row,
        column=2,
        value=value,
    )

    ws.cell(
        row=row,
        column=1,
    ).font = Font(
        bold=True,
    )

    ws.cell(
        row=row,
        column=1,
    ).fill = PatternFill(
        "solid",
        fgColor="D9EAF7",
    )

    ws.cell(
        row=row,
        column=1,
    ).border = BORDER

    ws.cell(
        row=row,
        column=2,
    ).border = BORDER


# ============================================================
# Summary
# ============================================================

def build_summary(
    wb: Workbook,
    site_code: str,
    results: dict[str, Any],
) -> None:

    ws = wb.active
    ws.title = "Summary"

    ws.merge_cells(
        "A1:F1"
    )

    ws["A1"] = (
        f"POP AUDIT REPORT - SITE {site_code}"
    )

    style_title(
        ws,
        "A1:F1",
    )

    audit_time = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    statuses = {
        name: service_status(result)
        for name, result in results.items()
    }

    overall = aggregate_status(
        list(
            statuses.values()
        )
    )

    add_kpi(
        ws,
        3,
        "Site Code",
        site_code,
    )

    add_kpi(
        ws,
        4,
        "Audit Time",
        audit_time,
    )

    add_kpi(
        ws,
        5,
        "Overall Status",
        overall,
    )

    style_status_cell(
        ws["B5"]
    )

    row = 7

    for service in (
        "cisco",
        "ericsson",
        "radio",
    ):

        add_kpi(
            ws,
            row,
            service.upper(),
            statuses.get(
                service,
                "NOT_AVAILABLE",
            ),
        )

        style_status_cell(
            ws.cell(
                row=row,
                column=2,
            )
        )

        row += 1

    ws.freeze_panes = "A3"

    auto_width(
        ws
    )


# ============================================================
# Devices
# ============================================================

def build_devices(
    wb: Workbook,
    results: dict[str, Any],
) -> None:

    ws = wb.create_sheet(
        "Devices"
    )

    headers = [
        "Service",
        "Hostname",
        "IP",
        "Zabbix Host ID",
        "Status",
    ]

    ws.append(
        headers
    )

    style_header(
        ws,
        1,
    )

    for service_name, result in results.items():

        data = find_data(
            result
        )

        if service_name == "cisco":

            for device in data.get(
                "results",
                [],
            ):

                ws.append([
                    "Cisco",
                    device.get("hostname"),
                    device.get("ip"),
                    device.get("zabbix_hostid"),
                    device.get("status"),
                ])

        elif service_name == "ericsson":

            for device in data.get(
                "results",
                [],
            ):

                ws.append([
                    "Ericsson",
                    device.get("hostname"),
                    device.get("ip"),
                    device.get("zabbix_hostid"),
                    device.get("status"),
                ])

        elif service_name == "radio":

            for link in data.get(
                "links",
                [],
            ):

                ws.append([
                    link.get("link_type"),
                    link.get("name"),
                    None,
                    link.get(
                        "zabbix_hostid"
                    ),
                    link.get("status"),
                ])

    style_body(
        ws,
        2,
        ws.max_row,
        ws.max_column,
    )

    for row in ws.iter_rows(
        min_row=2,
        min_col=5,
        max_col=5,
    ):
        style_status_cell(
            row[0]
        )

    ws.freeze_panes = "A2"

    auto_width(
        ws
    )


# ============================================================
# Cisco
# ============================================================

def build_cisco(
    wb: Workbook,
    result: dict[str, Any],
) -> None:

    ws = wb.create_sheet(
        "Cisco"
    )

    ws.append([
        "Hostname",
        "IP",
        "Interface",
        "Description",
        "Status",
        "Admin/Oper",
        "Speed",
        "Duplex",
        "CRC",
        "Input Errors",
        "Output Errors",
        "Output Drops",
        "Health Reason",
    ])

    style_header(
        ws,
        1,
    )

    data = find_data(
        result
    )

    for device in data.get(
        "results",
        [],
    ):

        checks = device.get(
            "checks",
            {}
        )

        interface_check = checks.get(
            "interfaces",
            {}
        )

        interfaces = interface_check.get(
            "data",
            []
        )

        if not interfaces:

            ws.append([
                device.get("hostname"),
                device.get("ip"),
                None,
                None,
                device.get("status"),
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
            ])

            continue

        for item in interfaces:

            ws.append([
                device.get("hostname"),
                device.get("ip"),
                item.get("interface"),
                item.get("description"),
                item.get("health"),
                (
                    f'{item.get("status_raw")} / '
                    f'{item.get("protocol") or ""}'
                ),
                item.get("speed"),
                item.get("duplex"),
                item.get("crc"),
                item.get("input_errors"),
                item.get("output_errors"),
                item.get("output_drops"),
                item.get("health_reason"),
            ])

    style_body(
        ws,
        2,
        ws.max_row,
        ws.max_column,
    )

    for row in ws.iter_rows(
        min_row=2,
        min_col=5,
        max_col=5,
    ):
        style_status_cell(
            row[0]
        )

    ws.freeze_panes = "A2"

    auto_width(
        ws
    )


# ============================================================
# Ericsson
# ============================================================

def build_ericsson(
    wb: Workbook,
    result: dict[str, Any],
) -> None:

    ws = wb.create_sheet(
        "Ericsson"
    )

    ws.append([
        "Hostname",
        "IP",
        "Overall Status",
        "LT All",
        "Cell Status",
        "Cell Count",
        "Alarm Status",
        "Alarm Count",
        "Critical",
        "Major",
        "Clock Status",
        "Clock State",
        "Clock Value",
    ])

    style_header(
        ws,
        1,
    )

    data = find_data(
        result
    )

    for device in data.get(
        "results",
        [],
    ):

        checks = device.get(
            "checks",
            {}
        )

        cells = checks.get(
            "cells",
            {}
        )

        alarms = checks.get(
            "alarms",
            {}
        )

        clock = checks.get(
            "clock",
            {}
        )

        ws.append([
            device.get("hostname"),
            device.get("ip"),
            device.get("status"),
            checks.get(
                "lt_all",
                {}
            ).get("status"),
            cells.get("status"),
            cells.get("count"),
            alarms.get("status"),
            alarms.get("total"),
            alarms.get("critical"),
            alarms.get("major"),
            clock.get("status"),
            clock.get("state"),
            clock.get("value"),
        ])

    style_body(
        ws,
        2,
        ws.max_row,
        ws.max_column,
    )

    for row in ws.iter_rows(
        min_row=2,
        max_row=ws.max_row,
    ):

        for index in (
            3,
            4,
            5,
            7,
            11,
        ):

            style_status_cell(
                row[index - 1]
            )

    ws.freeze_panes = "A2"

    auto_width(
        ws
    )

    # Alarm details
    detail = wb.create_sheet(
        "Ericsson Alarms"
    )

    detail.append([
        "Hostname",
        "Severity",
        "Timestamp",
        "Problem",
        "Additional Text",
        "Raw",
    ])

    style_header(
        detail,
        1,
    )

    for device in data.get(
        "results",
        [],
    ):

        alarms = (
            device.get(
                "checks",
                {}
            )
            .get(
                "alarms",
                {}
            )
            .get(
                "alarms",
                []
            )
        )

        for alarm in alarms:

            detail.append([
                device.get("hostname"),
                alarm.get("severity"),
                alarm.get("timestamp"),
                alarm.get("problem"),
                alarm.get("additional_text"),
                alarm.get("raw"),
            ])

    style_body(
        detail,
        2,
        detail.max_row,
        detail.max_column,
    )

    for row in detail.iter_rows(
        min_row=2,
        min_col=2,
        max_col=2,
    ):
        style_status_cell(
            row[0]
        )

    detail.freeze_panes = "A2"

    auto_width(
        detail,
        50,
    )


# ============================================================
# RTN
# ============================================================

def build_rtn(
    wb: Workbook,
    result: dict[str, Any],
) -> None:

    ws = wb.create_sheet(
        "RTN"
    )

    ws.append([
        "Radio Host",
        "Destination",
        "Overall Status",
        "Modulation Status",
        "Difference Status",
        "Traffic Status",
        "Packet Loss",
    ])

    style_header(
        ws,
        1,
    )

    data = find_data(
        result
    )

    for link in data.get(
        "links",
        [],
    ):

        if not isinstance(link, dict):
            continue

        if str(
            link.get(
                "link_type",
                ""
            )
        ).upper() != "RTN":

            continue

        analysis = link.get(
            "analysis",
            {}
        )

        if not isinstance(analysis, dict):
            continue

        analysis_links = analysis.get(
            "links",
            []
        )

        if not isinstance(analysis_links, list):
            continue

        for item in analysis_links:

            if not isinstance(item, dict):
                continue

            destination = item.get(
                "destination"
            )

            modulation = item.get(
                "modulation",
                []
            )

            difference = item.get(
                "difference",
                {}
            )

            traffic = item.get(
                "traffic",
                []
            )

            # ------------------------------------------------
            # Normalise modulation statuses.
            # Real Radio output may contain either:
            #   [{"status": "PASS"}, ...]
            # or:
            #   ["PASS", "WARNING", ...]
            # ------------------------------------------------
            modulation_statuses = []

            if isinstance(modulation, list):

                for value in modulation:

                    if isinstance(value, dict):

                        status = value.get(
                            "status"
                        )

                        if status:
                            modulation_statuses.append(
                                str(status).upper()
                            )

                    elif isinstance(value, str):

                        modulation_statuses.append(
                            value.upper()
                        )

            elif isinstance(modulation, dict):

                status = modulation.get(
                    "status"
                )

                if status:
                    modulation_statuses.append(
                        str(status).upper()
                    )

            elif isinstance(modulation, str):

                modulation_statuses.append(
                    modulation.upper()
                )

            # ------------------------------------------------
            # Normalise traffic statuses.
            # ------------------------------------------------
            traffic_statuses = []

            if isinstance(traffic, list):

                for value in traffic:

                    if isinstance(value, dict):

                        status = value.get(
                            "status"
                        )

                        if status:
                            traffic_statuses.append(
                                str(status).upper()
                            )

                    elif isinstance(value, str):

                        traffic_statuses.append(
                            value.upper()
                        )

            elif isinstance(traffic, dict):

                status = traffic.get(
                    "status"
                )

                if status:
                    traffic_statuses.append(
                        str(status).upper()
                    )

            elif isinstance(traffic, str):

                traffic_statuses.append(
                    traffic.upper()
                )

            # ------------------------------------------------
            # Difference may be dict or direct status/string.
            # ------------------------------------------------
            difference_status = None

            if isinstance(difference, dict):

                difference_status = difference.get(
                    "status"
                )

            elif isinstance(difference, str):

                difference_status = difference

            if difference_status:
                difference_status = str(
                    difference_status
                ).upper()

            # ------------------------------------------------
            # Packet-loss extraction.
            # ------------------------------------------------
            packet_loss = None

            packet_loss_data = item.get(
                "packet_loss"
            )

            if isinstance(packet_loss_data, dict):

                packet_loss = (
                    packet_loss_data.get("value")
                    if "value" in packet_loss_data
                    else packet_loss_data.get("status")
                )

            elif packet_loss_data is not None:

                packet_loss = packet_loss_data

            ws.append([
                link.get("name"),
                destination,
                item.get("status"),
                aggregate_status(
                    modulation_statuses
                ),
                difference_status,
                aggregate_status(
                    traffic_statuses
                ),
                packet_loss,
            ])

    style_body(
        ws,
        2,
        ws.max_row,
        ws.max_column,
    )

    for row in ws.iter_rows(
        min_row=2,
    ):

        for index in (
            3,
            4,
            5,
            6,
        ):

            style_status_cell(
                row[index - 1]
            )

    ws.freeze_panes = "A2"

    auto_width(
        ws
    )


# ============================================================
# EBand
# ============================================================

def build_eband(
    wb: Workbook,
    result: dict[str, Any],
) -> None:

    ws = wb.create_sheet(
        "EBand"
    )

    ws.append([
        "Radio Host",
        "Slot",
        "RX",
        "TX",
        "RX Calloff",
        "TX Calloff",
        "Calculated Difference",
        "Zabbix Difference",
        "Capacity",
        "RX Modulation",
        "RX Health",
        "Capacity Health",
        "Modulation Health",
        "Interface Status",
        "Overall Status",
    ])

    style_header(
        ws,
        1,
    )

    data = find_data(
        result
    )

    for link in data.get(
        "links",
        [],
    ):

        if str(
            link.get(
                "link_type",
                ""
            )
        ).upper() != "EBAND":

            continue

        analysis = link.get(
            "analysis",
            {}
        )

        checks = analysis.get(
            "checks",
            {}
        )

        difference = checks.get(
            "difference",
            {}
        )

        slots = difference.get(
            "slots",
            []
        )

        for slot_data in slots:

            rx = slot_data.get(
                "rx",
                {}
            )

            tx = slot_data.get(
                "tx",
                {}
            )

            rx_calloff = slot_data.get(
                "rx_calloff",
                {}
            )

            tx_calloff = slot_data.get(
                "tx_calloff",
                {}
            )

            diff = slot_data.get(
                "difference",
                {}
            )

            zabbix_diff = slot_data.get(
                "zabbix_difference",
                {}
            )

            capacity = slot_data.get(
                "capacity",
                {}
            )

            modulation = slot_data.get(
                "rx_modulation",
                {}
            )

            interfaces = checks.get(
                "interfaces",
                {}
            )

            ws.append([
                link.get("name"),
                slot_data.get("slot"),
                rx.get("value"),
                tx.get("value"),
                rx_calloff.get("value"),
                tx_calloff.get("value"),
                diff.get("value"),
                zabbix_diff.get("value"),
                capacity.get(
                    "metric",
                    {}
                ).get("value"),
                modulation.get(
                    "metric",
                    {}
                ).get("value"),
                slot_data.get(
                    "rx_health",
                    {}
                ).get("status"),
                capacity.get(
                    "health",
                    {}
                ).get("status"),
                modulation.get(
                    "health",
                    {}
                ).get("status"),
                interfaces.get(
                    "status"
                ),
                analysis.get(
                    "status"
                ),
            ])

    style_body(
        ws,
        2,
        ws.max_row,
        ws.max_column,
    )

    for row in ws.iter_rows(
        min_row=2,
    ):

        for index in (
            11,
            12,
            13,
            14,
            15,
        ):

            style_status_cell(
                row[index - 1]
            )

    ws.freeze_panes = "A2"

    auto_width(
        ws
    )


# ============================================================
# Workbook generation
# ============================================================

def create_workbook(
    site_code: str,
    results: dict[str, Any],
) -> tuple[Path, str]:

    wb = Workbook()

    build_summary(
        wb,
        site_code,
        results,
    )

    build_devices(
        wb,
        results,
    )

    if "cisco" in results:
        build_cisco(
            wb,
            results["cisco"],
        )

    if "ericsson" in results:
        build_ericsson(
            wb,
            results["ericsson"],
        )

    if "radio" in results:

        build_rtn(
            wb,
            results["radio"],
        )

        build_eband(
            wb,
            results["radio"],
        )

    timestamp = datetime.now().strftime(
        "%Y%m%d-%H%M%S"
    )

    safe_site = re.sub(
        r"[^A-Za-z0-9_-]",
        "_",
        site_code,
    )

    filename = (
        f"POP-Audit-{safe_site}-"
        f"{timestamp}.xlsx"
    )

    output = REPORT_DIR / filename

    wb.save(
        output
    )

    overall = aggregate_status(
        [
            service_status(result)
            for result in results.values()
        ]
    )

    return output, overall


# ============================================================
# API
# ============================================================

@app.get("/health")
def health():
    return {
        "service": "report",
        "status": "ok",
    }


@app.post("/generate")
def generate_report(
    payload: dict,
):

    site_code = str(
        payload.get(
            "site_code",
            "",
        )
    ).strip()

    if not site_code:
        raise HTTPException(
            status_code=400,
            detail="site_code is required",
        )

    results = payload.get(
        "results",
        {},
    )

    if not isinstance(
        results,
        dict,
    ):
        raise HTTPException(
            status_code=400,
            detail="results must be an object",
        )

    try:

        output, overall = create_workbook(
            site_code,
            results,
        )

    except Exception as exc:

        raise HTTPException(
            status_code=500,
            detail=(
                f"Report generation failed: {exc}"
            ),
        )

    return {
        "service": "report",
        "site_code": site_code,
        "status": overall,
        "filename": output.name,
        "path": str(output),
        "download_url": (
            f"/download/{output.name}"
        ),
    }


@app.get("/download/{filename}")
def download_report(
    filename: str,
):

    safe_name = Path(
        filename
    ).name

    path = REPORT_DIR / safe_name

    if not path.exists():
        raise HTTPException(
            status_code=404,
            detail="Report not found",
        )

    return FileResponse(
        path,
        filename=safe_name,
        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
